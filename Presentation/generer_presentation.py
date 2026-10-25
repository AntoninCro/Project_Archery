# -*- coding: utf-8 -*-
"""Génère la présentation orale de THUNK VR (PowerPoint 16:9), dans la direction artistique du jeu.

Fond violet de nuit, cadres dorés du pack Fantasy UI Borders de Kenney (les mêmes que les menus du jeu, teintés
avec les couleurs de `UiTheme`), titres en Cinzel avec le dégradé des titres du jeu (titleTop vers titleBottom).

    py Presentation/generer_presentation.py                    # THUNK_VR_Presentation.pptx (polices intégrées) et le PDF
    py Presentation/generer_presentation.py --apercus DOSSIER  # en plus, une image PNG par diapo

Les captures du jeu se lisent dans `Presentation/captures/` (titre.png, tir.png… : la liste est dans CAPTURES_ATTENDUES).
Une capture absente laisse un cadre « capture à venir » : relancer le script après avoir ajouté les captures.
La fin (polices intégrées au fichier, PDF de secours, aperçus) passe par PowerPoint, qui doit être installé, avec la
police Cinzel (dans `Project Archery/Assets/cinzel`).
"""
import argparse
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image, ImageChops, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ICI = Path(__file__).resolve().parent
ASSETS = ICI.parent / "Project Archery" / "Assets"
KENNEY = ASSETS / "kenney_fantasy-ui-borders" / "PNG" / "Double"
ICONES = ASSETS / "kenney_game-icons" / "PNG" / "White" / "2x"
CAPTURES = ICI / "captures"
SORTIE = ICI / "THUNK_VR_Presentation.pptx"

CAPTURES_ATTENDUES = {
    "titre": "un plan large : la clairière et la tour, ou la porte de l'antre",
    "tir": "l'arc tendu, avec l'anneau de timing",
    "vague": "des ennemis qui arrivent sur une voie",
    "boutique": "la boutique entre deux vagues",
    "mk": "MK, le dragon, dans son antre",
    "construction": "une baliste ou un plan bleu, avec le marteau",
    "armes": "le menu des armes",
    "grenade": "la grenade à flèches",
    "secret": "le portail des étoiles, le toit cyberpunk ou une faille du katana",
    "unity": "l'éditeur Unity, la scène ouverte",
    "github": "le graphe des branches sur GitHub",
}

# Diapo 16:9, en pouces.
L, H = 13.333, 7.5
MARGE = 0.6
DPI = 300          # résolution des cadres générés
ECHELLE_CADRE = 2  # taille des pixels des bordures Kenney (lignes nettes)

# Couleurs de UiTheme (Resources/UiTheme.asset).
FOND = "0E0920"
PANNEAU = "170F30"      # panelColor
CARTE = "1F1740"        # cardColor
OR = "FFCC61"           # panelBorderColor
TITRE_HAUT = "FFF5B8"   # titleTop
TITRE_BAS = "FFA333"    # titleBottom
TEXTE = "F5EDDB"        # textColor
LABEL = "D1DBFF"        # labelColor
DISCRET = "A79FC4"
DANGER = "E0475E"       # dangerButton, éclairci pour le fond sombre
VERT = "66F28C"         # sliderFills[3]

# Une couleur par orateur (les remplissages des curseurs du jeu).
ORATEURS = {
    "Merleau": "4DD9FF",
    "Jude": "F259D9",
    "Antonin": "FFBD40",
    "Dovin": VERT,
}

# Cinzel toujours en graisse normale : c'est la seule que PowerPoint intègre au fichier (les .otf du jeu).
TITRE = dict(police="Cinzel Decorative", taille=32, degrade=(TITRE_HAUT, TITRE_BAS), ombre=True)
ETIQUETTE = dict(police="Cinzel", taille=13, couleur=OR, espacement=120)
CORPS = dict(police="Calibri", taille=16, couleur=TEXTE)
LEGENDE = dict(police="Calibri", taille=12, couleur=DISCRET)
# Les chiffres : le « 1 » de Cinzel ressemble à un « I » (« 10/10 » se lisait « IO/IO »).
CHIFFRE = dict(police="Cambria", taille=40, gras=True, degrade=(TITRE_HAUT, TITRE_BAS), ombre=True)

A = "http://schemas.openxmlformats.org/drawingml/2006/main"

TEMP = Path(tempfile.mkdtemp(prefix="thunk_ppt_"))
_compteur = [0]


def _fichier(suffixe):
    _compteur[0] += 1
    return TEMP / f"img{_compteur[0]:03d}{suffixe}"


def rgb(hexa):
    return RGBColor.from_string(hexa)


# ---------------------------------------------------------------------------------------------------------------
# Images générées : fond, cadres Kenney, voiles, icônes.

def _sprite(nom):
    return Image.open(KENNEY / nom).convert("RGBA")


def _teinte(im, couleur, alpha=255):
    r, g, b = bytes.fromhex(couleur)
    masque = im.split()[3].point(lambda v: v * alpha // 255)
    out = Image.new("RGBA", im.size, (r, g, b, 0))
    out.putalpha(masque)
    return out


def _neuf_parts(im, bord, largeur, hauteur, echelle):
    """Découpe en 9 (comme le Sliced de Unity) : coins à l'échelle, bords et centre étirés, pixels nets."""
    b, s = bord, echelle
    bs = b * s
    w0, h0 = im.size
    cw, ch = max(1, largeur - 2 * bs), max(1, hauteur - 2 * bs)
    out = Image.new("RGBA", (largeur, hauteur), (0, 0, 0, 0))

    def part(boite, taille):
        return im.crop(boite).resize(taille, Image.NEAREST)

    out.paste(part((0, 0, b, b), (bs, bs)), (0, 0))
    out.paste(part((w0 - b, 0, w0, b), (bs, bs)), (largeur - bs, 0))
    out.paste(part((0, h0 - b, b, h0), (bs, bs)), (0, hauteur - bs))
    out.paste(part((w0 - b, h0 - b, w0, h0), (bs, bs)), (largeur - bs, hauteur - bs))
    out.paste(part((b, 0, w0 - b, b), (cw, bs)), (bs, 0))
    out.paste(part((b, h0 - b, w0 - b, h0), (cw, bs)), (bs, hauteur - bs))
    out.paste(part((0, b, b, h0 - b), (bs, ch)), (0, bs))
    out.paste(part((w0 - b, b, w0, h0 - b), (bs, ch)), (largeur - bs, bs))
    out.paste(part((b, b, w0 - b, h0 - b), (cw, ch)), (bs, bs))
    return out


# Les trois cadres du jeu (UiTheme) : panneau, cadre, bouton ; remplissage et bordure, avec leurs découpes.
STYLES = {
    "panneau": ("Panel/panel-010.png", 24, "Border/panel-border-010.png", 20),
    "cadre": ("Panel/panel-008.png", 32, "Border/panel-border-008.png", 16),
    "bouton": ("Panel/panel-002.png", 28, "Border/panel-border-002.png", 24),
}


def image_cadre(w, h, style="panneau", fond=PANNEAU, alpha=240, bordure=OR):
    remplissage, b_fond, contour, b_contour = STYLES[style]
    pw, ph = round(w * DPI), round(h * DPI)
    img = Image.new("RGBA", (pw, ph), (0, 0, 0, 0))
    if fond:
        img.alpha_composite(_neuf_parts(_teinte(_sprite(remplissage), fond, alpha), b_fond, pw, ph, ECHELLE_CADRE))
    if bordure:
        img.alpha_composite(_neuf_parts(_teinte(_sprite(contour), bordure), b_contour, pw, ph, ECHELLE_CADRE))
    chemin = _fichier(".png")
    img.save(chemin, optimize=True)
    return chemin


def image_fond():
    """Nuit violette : une lueur douce en haut, les bords plus sombres."""
    w, h = 1600, 900
    base = Image.new("RGB", (w, h), tuple(bytes.fromhex(FOND)))
    lueur = Image.new("L", (w, h), 0)
    d = ImageDraw.Draw(lueur)
    for i in range(60, 0, -1):
        t = i / 60
        rx, ry = 1250 * t, 620 * t
        d.ellipse((w * 0.62 - rx, -260 - ry, w * 0.62 + rx, -260 + ry + 260), fill=int(150 * (1 - t) ** 1.6))
    teinte = Image.new("RGB", (w, h), tuple(bytes.fromhex("3A2470")))
    base = Image.composite(teinte, base, lueur)
    ombre = Image.new("L", (w, h), 120)
    ImageDraw.Draw(ombre).rounded_rectangle((140, 90, w - 140, h - 90), radius=160, fill=0)
    ombre = ombre.filter(ImageFilter.GaussianBlur(120))
    noir = Image.new("RGB", (w, h), (4, 2, 10))
    base = Image.composite(noir, base, ombre)
    chemin = _fichier(".png")
    base.save(chemin, optimize=True)
    return chemin


def image_voile(gauche=0.8, bas=0.0):
    """Voile sombre par-dessus une capture plein écran, pour que le texte se lise : opaque sur la gauche (le texte),
    qui s'efface jusqu'à `gauche` (fraction de la largeur), et en bas sur la hauteur `bas`."""
    w, h = 1600, 900
    r, g, b = bytes.fromhex(FOND)

    def lisse(t):
        t = min(1.0, max(0.0, t))
        return t * t * (3 - 2 * t)

    colonne = Image.new("L", (w, 1))
    for x in range(w):
        colonne.putpixel((x, 0), int(55 + 190 * (1 - lisse((x / w - 0.42) / (gauche - 0.42)))))
    alpha = colonne.resize((w, h))
    if bas:
        ligne = Image.new("L", (1, h))
        for y in range(h):
            ligne.putpixel((0, y), int(55 + 185 * lisse((y / h - (1 - bas)) / bas)))
        alpha = ImageChops.lighter(alpha, ligne.resize((w, h)))
    img = Image.new("RGBA", (w, h), (r, g, b, 0))
    img.putalpha(alpha)
    chemin = _fichier(".png")
    img.save(chemin, optimize=True)
    return chemin


def image_icone(nom, couleur):
    im = Image.open(ICONES / f"{nom}.png").convert("RGBA")
    chemin = _fichier(".png")
    _teinte(im, couleur).save(chemin)
    return chemin


def chercher_capture(nom):
    for ext in (".png", ".jpg", ".jpeg", ".webp", ".PNG", ".JPG"):
        f = CAPTURES / f"{nom}{ext}"
        if f.exists():
            return f
    return None


def image_capture(source, w, h):
    """Recadre la capture au centre, au format du cadre, et l'allège (JPEG)."""
    im = Image.open(source).convert("RGB")
    cible = w / h
    sw, sh = im.size
    if sw / sh > cible:
        nw = round(sh * cible)
        im = im.crop(((sw - nw) // 2, 0, (sw - nw) // 2 + nw, sh))
    else:
        nh = round(sw / cible)
        im = im.crop((0, (sh - nh) // 2, sw, (sh - nh) // 2 + nh))
    largeur = min(im.width, round(w * 160))
    im = im.resize((largeur, round(largeur / cible)), Image.LANCZOS)
    chemin = _fichier(".jpg")
    im.save(chemin, quality=88)
    return chemin


# ---------------------------------------------------------------------------------------------------------------
# Texte : styles, dégradé des titres, puces.

def _degrade(rPr, haut, bas):
    for tag in ("a:solidFill", "a:gradFill", "a:noFill"):
        for el in rPr.findall(qn(tag)):
            rPr.remove(el)
    rPr.insert(0, parse_xml(
        f'<a:gradFill xmlns:a="{A}" rotWithShape="1"><a:gsLst>'
        f'<a:gs pos="0"><a:srgbClr val="{haut}"/></a:gs><a:gs pos="100000"><a:srgbClr val="{bas}"/></a:gs>'
        f'</a:gsLst><a:lin ang="5400000" scaled="0"/></a:gradFill>'))


def _ombre(rPr):
    effet = parse_xml(
        f'<a:effectLst xmlns:a="{A}"><a:outerShdw blurRad="38100" dist="25400" dir="5400000" algn="t" '
        f'rotWithShape="0"><a:srgbClr val="000000"><a:alpha val="75000"/></a:srgbClr></a:outerShdw></a:effectLst>')
    remplissage = rPr.find(qn("a:gradFill"))
    if remplissage is None:
        remplissage = rPr.find(qn("a:solidFill"))
    if remplissage is not None:
        remplissage.addnext(effet)
    else:
        rPr.insert(0, effet)


def _style(run, st):
    f = run.font
    f.name = st.get("police", "Calibri")
    f.size = Pt(st.get("taille", 16))
    f.bold = st.get("gras", False)
    f.italic = st.get("italique", False)
    rPr = run._r.get_or_add_rPr()
    if "espacement" in st:
        rPr.set("spc", str(st["espacement"]))
    if st.get("degrade"):
        _degrade(rPr, *st["degrade"])
    else:
        f.color.rgb = rgb(st.get("couleur", TEXTE))
    if st.get("ombre"):
        _ombre(rPr)


def _puce(p, couleur=OR, retrait=0.26):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(retrait))))
    pPr.set("indent", str(-int(Inches(retrait))))
    for xml in (f'<a:buClr xmlns:a="{A}"><a:srgbClr val="{couleur}"/></a:buClr>',
                f'<a:buSzPct xmlns:a="{A}" val="70000"/>',
                f'<a:buFont xmlns:a="{A}" typeface="Segoe UI Symbol"/>',
                f'<a:buChar xmlns:a="{A}" char="◆"/>'):
        pPr.append(parse_xml(xml))


def typo(s):
    """Espaces insécables du français : jamais de « ; » ou de « : » en début de ligne."""
    for a, b in ((" ;", "\u00a0;"), (" :", "\u00a0:"), (" !", "\u00a0!"), (" ?", "\u00a0?"),
                 ("« ", "«\u00a0"), (" »", "\u00a0»")):
        s = s.replace(a, b)
    return re.sub(r"(\d) (?=\w)", "\\1\u00a0", s)  # « 72 images/s » reste sur une ligne


def liste(*elements):
    """Des éléments séparés par « · », coupés seulement entre deux éléments."""
    return " · ".join(e.replace(" ", " ") for e in elements)


def texte(diapo, x, y, w, h, paragraphes, ancre="haut", nom=None):
    """paragraphes : liste de dicts {runs: [(texte, style)], align, apres, avant, interligne, puce}
    ou de tuples (texte, style) pour un paragraphe d'un seul morceau."""
    boite = diapo.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if nom:
        boite.name = nom
    tf = boite.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"haut": MSO_ANCHOR.TOP, "milieu": MSO_ANCHOR.MIDDLE, "bas": MSO_ANCHOR.BOTTOM}[ancre]
    for i, para in enumerate(paragraphes):
        if isinstance(para, tuple):
            para = {"runs": [para]}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"g": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "d": PP_ALIGN.RIGHT}[para.get("align", "g")]
        if "apres" in para:
            p.space_after = Pt(para["apres"])
        if "avant" in para:
            p.space_before = Pt(para["avant"])
        if "interligne" in para:
            p.line_spacing = para["interligne"]
        if para.get("puce"):
            _puce(p, para.get("couleur_puce", OR))
        for contenu, style in para["runs"]:
            run = p.add_run()
            run.text = typo(contenu)
            _style(run, style)
    return boite


def st(base, **modifs):
    d = dict(base)
    d.update(modifs)
    return d


# ---------------------------------------------------------------------------------------------------------------
# Éléments de diapo.

def image(diapo, chemin, x, y, w, h, nom=None):
    pic = diapo.shapes.add_picture(str(chemin), Inches(x), Inches(y), Inches(w), Inches(h))
    if nom:
        pic.name = nom
    return pic


def cadre(diapo, x, y, w, h, style="panneau", fond=PANNEAU, alpha=240, bordure=OR, nom=None):
    return image(diapo, image_cadre(w, h, style, fond, alpha, bordure), x, y, w, h, nom)


def capture(diapo, nom, x, y, w, h):
    """Une capture du jeu dans le cadre doré des menus ; un cadre vide si elle manque."""
    marge = 0.07
    source = chercher_capture(nom)
    if source:
        cadre(diapo, x - marge, y - marge, w + 2 * marge, h + 2 * marge, "cadre", fond=FOND, alpha=255,
              bordure=None, nom=f"Fond {nom}")
        image(diapo, image_capture(source, w, h), x, y, w, h, nom=f"Capture {nom}")
    else:
        cadre(diapo, x - marge, y - marge, w + 2 * marge, h + 2 * marge, "cadre", fond=CARTE, alpha=255,
              bordure=None, nom=f"Fond {nom}")
        texte(diapo, x + 0.3, y, w - 0.6, h, [
            {"runs": [("CAPTURE À VENIR", st(ETIQUETTE, taille=11, couleur=DISCRET))], "align": "c", "apres": 4},
            {"runs": [(f"captures/{nom}.png", st(LEGENDE, couleur=LABEL))], "align": "c", "apres": 2},
            {"runs": [(CAPTURES_ATTENDUES[nom], LEGENDE)], "align": "c"},
        ], ancre="milieu", nom=f"Attente {nom}")
    cadre(diapo, x - marge, y - marge, w + 2 * marge, h + 2 * marge, "cadre", fond=None, nom=f"Cadre {nom}")


def capture_plein_ecran(diapo, nom, voile_gauche=0.8, voile_bas=0.0):
    source = chercher_capture(nom)
    if not source:
        return False
    image(diapo, image_capture(source, L, H), 0, 0, L, H, nom=f"Capture {nom}")
    image(diapo, image_voile(voile_gauche, voile_bas), 0, 0, L, H, nom="Voile")
    return True


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(L)
        self.prs.slide_height = Inches(H)
        self.vide = self.prs.slide_layouts[6]
        self.fond = image_fond()
        self.total = 12
        self.numero = 0

    def diapo(self, titre=None, orateur=None, fond=True):
        self.numero += 1
        d = self.prs.slides.add_slide(self.vide)
        if fond:
            image(d, self.fond, 0, 0, L, H, nom="Fond")
        if titre:
            texte(d, MARGE, 0.45, 10.2, 0.8, [{"runs": [(titre, TITRE)]}], ancre="milieu", nom="Titre")
        if orateur:
            texte(d, 10.4, 0.62, L - MARGE - 10.4, 0.45, [{"runs": [
                (orateur.upper(), st(ETIQUETTE, taille=14, couleur=ORATEURS[orateur], espacement=250))],
                "align": "d"}], ancre="milieu", nom="Orateur")
        if self.numero > 1:
            texte(d, MARGE, 7.0, 3, 0.3, [("THUNK VR", st(ETIQUETTE, taille=10, couleur=DISCRET, espacement=200))],
                  nom="Pied")
            texte(d, L - MARGE - 1.5, 7.0, 1.5, 0.3, [{"runs": [
                (f"{self.numero} / {self.total}", st(LEGENDE, taille=11))], "align": "d"}],
                nom="Numéro")
        return d


def notes(diapo, contenu):
    diapo.notes_slide.notes_text_frame.text = contenu


# ---------------------------------------------------------------------------------------------------------------
# Les diapos.

def d01_titre(deck):
    d = deck.diapo(fond=True)
    plein = capture_plein_ecran(d, "titre", voile_gauche=0.85, voile_bas=0.35)
    texte(d, 0.85, 2.0, 9, 0.4, [("UN JEU DE TIR À L'ARC EN RÉALITÉ VIRTUELLE",
                                  st(ETIQUETTE, taille=14, espacement=300))], nom="Accroche")
    texte(d, 0.8, 2.4, 10, 1.5, [("THUNK VR", st(TITRE, police="Cinzel Decorative", taille=88))],
          ancre="milieu", nom="Nom du jeu")
    texte(d, 0.85, 3.95, 9, 0.5, [("Survie, roguelike et tower defense, sur Meta Quest 3",
                                   st(CORPS, taille=22))], nom="Sous-titre")
    noms = [("Antonin Croquette", "Antonin"), ("Jude Aybalen", "Jude"),
            ("Tchoffo Manfouo Dovin Vilson", "Dovin"), ("Merleau Gezo", "Merleau")]
    runs = []
    for i, (n, o) in enumerate(noms):
        if i:
            runs.append(("   ·   ", st(CORPS, taille=16, couleur=DISCRET)))
        runs.append((n, st(CORPS, taille=16, gras=True, couleur=ORATEURS[o])))
    texte(d, 0.85, 5.75, 11.5, 0.4, [{"runs": runs}], nom="Équipe")
    texte(d, 0.85, 6.2, 11, 0.35, [("ST2OOS · Unity XR · 23 octobre 2026", st(LEGENDE, taille=13))], nom="Cours")
    if not plein:
        capture(d, "titre", 9.55, 1.6, 3.18, 3.18 * 9 / 16)
    notes(d, "MERLEAU — environ 20 s.\n"
             "Bonjour. Nous sommes Antonin, Jude, Dovin et moi, Merleau. On va vous présenter THUNK VR, notre jeu "
             "de tir à l'arc en réalité virtuelle pour le Meta Quest 3. Le nom vient du bruit d'une flèche qui se "
             "plante : « thunk ».")


def d02_programme(deck):
    d = deck.diapo("Au programme", "Merleau")
    parties = [
        ("Merleau", "target", "Le concept", "Le jeu, ses trois genres, le déroulement d'une partie"),
        ("Jude", "checkmark", "Ce qu'on a réalisé", "Les consignes du cours, et tout ce qu'on a ajouté"),
        ("Antonin", "gear", "La technique", "Comment le jeu est construit, et ce qui nous a résisté"),
        ("Dovin", "multiplayer", "L'équipe", "Qui a fait quoi, comment on a travaillé, le bilan"),
    ]
    w, gap, y, h = (L - 2 * MARGE - 3 * 0.3) / 4, 0.3, 1.65, 3.55
    for i, (o, icone, sujet, detail) in enumerate(parties):
        x = MARGE + i * (w + gap)
        cadre(d, x, y, w, h, nom=f"Partie {i + 1}")
        image(d, image_icone(icone, ORATEURS[o]), x + 0.38, y + 0.42, 0.5, 0.5, nom=f"Icône partie {i + 1}")
        texte(d, x + 0.35, y + 1.15, w - 0.7, h - 1.5, [
            {"runs": [(o.upper(), st(ETIQUETTE, taille=12, couleur=ORATEURS[o], espacement=250))], "apres": 10},
            {"runs": [(sujet, st(CORPS, taille=20, gras=True))], "apres": 6},
            {"runs": [(detail, st(CORPS, taille=15, couleur=LABEL))]},
        ], nom=f"Texte partie {i + 1}")
        texte(d, x + 0.35, y + h - 0.65, w - 0.7, 0.35, [("2 MIN 30", st(ETIQUETTE, taille=11, couleur=DISCRET))],
              nom=f"Durée partie {i + 1}")
    bw = 7.2
    cadre(d, (L - bw) / 2, 5.6, bw, 0.95, "bouton", fond=CARTE, nom="Démo et questions")
    texte(d, (L - bw) / 2, 5.6, bw, 0.95, [{"runs": [
        ("Puis ", st(CORPS, taille=17)), ("5 minutes de démo au casque", st(CORPS, taille=17, gras=True, couleur=OR)),
        (", et vos questions", st(CORPS, taille=17))], "align": "c"}], ancre="milieu", nom="Texte démo")
    notes(d, "MERLEAU — environ 15 s.\n"
             "On s'est partagé la présentation en quatre parties de deux minutes et demie : je commence par le "
             "concept, Jude présente ce qu'on a réalisé, Antonin la technique et les difficultés, et Dovin "
             "l'organisation de l'équipe et le bilan. Ensuite, cinq minutes de démo au casque, puis vos questions.")


def d03_concept(deck):
    d = deck.diapo("Défendre sa tour, à l'arc", "Merleau")
    cw = 7.0
    ch = cw * 9 / 16
    cy = 1.7
    capture(d, "tir", MARGE, cy, cw, ch)
    texte(d, MARGE, cy + ch + 0.2, cw, 0.5, [(
        "Tendre la corde : l'anneau passe du rouge au vert. Lâcher dans le vert, c'est le tir parfait.",
        LEGENDE)], nom="Légende tir")
    px, pw = MARGE + cw + 0.45, L - 2 * MARGE - cw - 0.45
    cadre(d, px, cy - 0.07, pw, ch + 0.14, nom="Trois genres")
    genres = [
        ("SURVIE", "Six vagues de monstres de plus en plus fortes, chacune finie par un boss."),
        ("ROGUELIKE", "Entre les vagues, des améliorations tirées au sort, qui se cumulent sans limite."),
        ("TOWER DEFENSE", "Récolter des cristaux, puis construire ses défenses au marteau, avec ses mains."),
    ]
    paras = []
    for i, (g, t) in enumerate(genres):
        paras.append({"runs": [(g, st(ETIQUETTE, espacement=200))], "apres": 3, "avant": 0 if i == 0 else 14})
        paras.append({"runs": [(t, st(CORPS, taille=16))]})
    texte(d, px + 0.4, cy + 0.3, pw - 0.8, ch - 0.6, paras, ancre="milieu", nom="Texte genres")
    texte(d, px, cy + ch + 0.2, pw, 0.5, [{"runs": [
        ("Inspirations : ", LEGENDE), ("Megabonk", st(LEGENDE, italique=True, couleur=LABEL)),
        (", et ", LEGENDE), ("Longbow", st(LEGENDE, italique=True, couleur=LABEL)),
        (" dans ", LEGENDE), ("The Lab", st(LEGENDE, italique=True, couleur=LABEL))]}], nom="Inspirations")
    notes(d, "MERLEAU — environ 1 min.\n"
             "THUNK VR mélange trois genres. De la survie : on tient face à six vagues de monstres de plus en plus "
             "fortes, et chacune se termine par un boss. Du roguelike : entre les vagues, on achète des améliorations "
             "tirées au sort, qui se cumulent sans limite ; chaque partie est différente. Et du tower defense : "
             "pendant les vagues, on sort récolter des cristaux et on construit ses défenses au marteau, avec ses "
             "mains.\n"
             "Au cœur du jeu, le tir à l'arc est physique : on prend la flèche dans son dos, on l'encoche, on tend la "
             "corde. Un anneau de timing récompense le tir parfait, lâché dans le vert.\n"
             "Nos inspirations : Megabonk pour le mélange survie et roguelike, et Longbow, dans The Lab de Valve, "
             "pour la défense à l'arc en VR.")


def d04_partie(deck):
    d = deck.diapo("Une partie, de la tour au dragon", "Merleau")
    etapes = [
        ("MENU", "On choisit la difficulté : le ciel change."),
        ("VAGUE", "Une accalmie pour construire, puis la montée."),
        ("BOSS", "Un boss par vague, avec son point faible."),
        ("PAUSE", "Boutique et coffres ; le gong relance."),
        ("MK", "Vague 6 : le dragon, dans son antre."),
    ]
    gap = 0.42
    w = (L - 2 * MARGE - 4 * gap) / 5
    y, h = 1.6, 1.6
    fleche = image_icone("arrowRight", OR)
    for i, (nom, t) in enumerate(etapes):
        x = MARGE + i * (w + gap)
        cadre(d, x, y, w, h, "bouton", fond=CARTE, nom=f"Étape {i + 1}")
        texte(d, x + 0.28, y + 0.25, w - 0.56, h - 0.45, [
            {"runs": [(nom, st(ETIQUETTE, espacement=200))], "apres": 6},
            {"runs": [(t, st(CORPS, taille=14))]},
        ], nom=f"Texte étape {i + 1}")
        if i < 4:
            image(d, fleche, x + w + (gap - 0.24) / 2, y + h / 2 - 0.12, 0.24, 0.24, nom=f"Flèche {i + 1}")
    cw = (L - 2 * MARGE - 2 * 0.35) / 3
    ch = cw * 9 / 16
    cy = 3.65
    legendes = [("vague", "Une vague arrive par la voie du milieu"), ("boutique", "La boutique, entre deux vagues"),
                ("mk", "MK, le boss final")]
    for i, (n, leg) in enumerate(legendes):
        x = MARGE + i * (cw + 0.35)
        capture(d, n, x, cy, cw, ch)
        texte(d, x, cy + ch + 0.15, cw, 0.3, [(leg, LEGENDE)], nom=f"Légende {n}")
    notes(d, "MERLEAU — environ 1 min.\n"
             "Une partie commence en haut de la tour. On choisit la difficulté sur des panneaux, et le ciel change : "
             "jour, coucher de soleil, nuit, lune de sang. Un coup de gong lance la vague : d'abord une accalmie, "
             "pour récolter et construire, puis la montée, où les ennemis sortent de la forêt de plus en plus vite. "
             "À la fin du chrono arrive le boss, avec un point faible qui bat comme un cœur.\n"
             "Entre les vagues : la boutique et les coffres ; on tire dans le gong pour repartir. À la vague 6, on "
             "descend affronter MK, un dragon, dans son antre. Après la victoire, on continue en mode infini, ou on "
             "inscrit son score au classement. Une partie complète dure environ 25 minutes.\n"
             "Je laisse la parole à Jude.")


def d05_consignes(deck):
    d = deck.diapo("Toutes les consignes sont dans le jeu", "Jude")
    consignes = [
        ("Interactions XR naturelles", "Prendre la flèche dans son dos, l'encocher, tendre, lâcher."),
        ("Trajectoire physique", "La gravité ; la portée dépend de la tension et du timing."),
        ("Plusieurs arcs et flèches", "8 armes ; flèches de foudre, de glace, explosives, ricochet…"),
        ("Améliorations entre les niveaux", "Une boutique entre les vagues : 23 améliorations, 3 raretés."),
        ("Zones de score, score en direct", "Tête, corps, points faibles ; score et combo sur la montre."),
        ("Niveaux de difficulté", "Facile, Normal, Difficile, Impossible, chacun avec son ciel."),
        ("Mode entraînement", "Des cibles, dont une mobile, au menu et pendant les pauses."),
        ("Menus et paramètres", "Menus dans le décor, tutoriel en vidéos, réglages de confort."),
        ("Clavier, sauvegarde, classement", "Nom au clavier virtuel d'XRI ; classement en JSON."),
        ("Sons 3D", "Sons spatialisés, ambiances de jour et de nuit, musiques."),
    ]
    coche = image_icone("checkmark", OR)
    pw = (L - 2 * MARGE - 0.35) / 2
    py, ph = 1.55, 5.2
    rh = (ph - 0.5) / 5
    for c in range(2):
        px = MARGE + c * (pw + 0.35)
        cadre(d, px, py, pw, ph, nom=f"Colonne {c + 1}")
        for r in range(5):
            label, t = consignes[c * 5 + r]
            y = py + 0.27 + r * rh
            image(d, coche, px + 0.35, y + 0.12, 0.26, 0.26, nom=f"Coche {c * 5 + r + 1}")
            texte(d, px + 0.85, y + 0.06, pw - 1.2, rh - 0.1, [
                {"runs": [(label.upper(), st(ETIQUETTE, taille=12, espacement=120))], "apres": 3},
                {"runs": [(t, st(CORPS, taille=15))]},
            ], nom=f"Consigne {c * 5 + r + 1}")
    notes(d, "JUDE — environ 1 min 15.\n"
             "Le cours demandait une liste de fonctionnalités : elles sont toutes dans le jeu. Les interactions "
             "passent par des gestes naturels : prendre la flèche dans son dos, tendre la corde, lâcher. La flèche "
             "suit une vraie trajectoire, avec la gravité, et sa portée dépend de la tension et du timing.\n"
             "Il y a plusieurs arcs et des flèches spéciales, une boutique d'améliorations entre les vagues, des "
             "zones de score : la tête et le corps des ennemis, les points faibles des boss. Le score s'affiche en "
             "direct sur une montre au poignet.\n"
             "Enfin : quatre difficultés, un mode entraînement avec des cibles, des menus dans le décor, un clavier "
             "virtuel pour saisir son nom, le classement sauvegardé en JSON, et des sons en 3D.")


def d06_ajouts(deck):
    d = deck.diapo("Et tout ce qu'on a ajouté", "Jude")
    ajouts = [
        ("construction", "Construire au marteau", [
            ("12 socles : barricades, piques, tours de givre et balistes, à 3 niveaux.", None)]),
        ("armes", "Le menu des armes", [
            ("Les arcs deviennent des classes, débloquées par des succès. ", None), ("Jude", "Jude")]),
        ("grenade", "La grenade à flèches", [
            ("Sortie de la poche, elle éclate en 28 flèches. ", None), ("Dovin", "Dovin")]),
        ("secret", "Deux secrets", [
            ("Le portail des étoiles et le toit cyberpunk, chacun avec son arme cachée.", None)]),
    ]
    gap = 0.32
    cw = (L - 2 * MARGE - 3 * gap) / 4
    ch = cw * 9 / 16
    for i, (n, titre, desc) in enumerate(ajouts):
        x = MARGE + i * (cw + gap)
        capture(d, n, x, 1.6, cw, ch)
        runs = [(t, st(CORPS, taille=15) if o is None else st(CORPS, taille=15, gras=True, couleur=ORATEURS[o]))
                for t, o in desc]
        texte(d, x, 1.6 + ch + 0.28, cw, 1.6, [
            {"runs": [(titre.upper(), st(ETIQUETTE, taille=12, espacement=150))], "apres": 5},
            {"runs": runs},
        ], nom=f"Ajout {i + 1}")
    by, bh = 5.3, 1.35
    cadre(d, MARGE, by, L - 2 * MARGE, bh, nom="Aussi")
    texte(d, MARGE + 0.45, by, L - 2 * MARGE - 0.9, bh, [
        {"runs": [("AUSSI", st(ETIQUETTE, espacement=200))], "apres": 4},
        {"runs": [(liste("cinq boss, et MK, un dragon en trois phases", "courir en balançant les bras",
                         "glisser et sauter", "des coffres", "un mode infini"),
                   st(CORPS, taille=16))]},
    ], ancre="milieu", nom="Texte aussi")
    notes(d, "JUDE — environ 1 min 15.\n"
             "On est aussi allés plus loin. Le tower defense : douze socles où construire barricades, pièges, tours "
             "de givre et balistes, en frappant un plan au marteau.\n"
             "Ma partie, c'est le menu des armes au début de la partie : les arcs sont devenus des classes, chacune "
             "avec ses caractéristiques, débloquées par des succès. J'ai ajouté le système de succès et les arcs "
             "court, puissant et runique.\n"
             "Dovin a fait la grenade à flèches : on la sort de sa poche, on la lance, et elle éclate en 28 flèches. "
             "Et deux secrets sont cachés sur la carte : le portail des étoiles et le toit cyberpunk, chacun avec "
             "une arme cachée. Il y a aussi cinq boss et le dragon final, la course en balançant les bras, et un mode "
             "infini.\n"
             "Antonin va vous expliquer comment tout ça est construit.")


def d07_technique(deck):
    d = deck.diapo("Comment c'est construit", "Antonin")
    gw = 6.35
    texte(d, MARGE, 1.55, gw, 0.95, [
        {"runs": [("LE SOCLE", st(ETIQUETTE, espacement=200))], "apres": 4},
        {"runs": [("Unity 6 et URP · XR Interaction Toolkit 3.5 · OpenXR · Meta Quest 3 relié au PC par Link",
                   st(CORPS, taille=15))]},
    ], nom="Socle")
    choix = [
        ("Les réglages en ScriptableObjects", " : on équilibre sans toucher au code."),
        ("Des pools d'objets", " : flèches, ennemis et effets réutilisés, pour tenir 72 images/s."),
        ("Le NavMesh", " : les ennemis suivent les trois voies jusqu'à la tour."),
        ("Un mode clavier-souris", " : on teste sans mettre le casque."),
    ]
    paras = [{"runs": [("NOS CHOIX", st(ETIQUETTE, espacement=200))], "apres": 4}]
    for gras, suite in choix:
        paras.append({"runs": [(gras, st(CORPS, taille=15, gras=True)), (suite, st(CORPS, taille=15))],
                      "apres": 5, "puce": True})
    texte(d, MARGE, 2.65, gw, 2.3, paras, nom="Choix")
    iy, ih = 5.0, 1.65
    cadre(d, MARGE, iy, gw, ih, "panneau", fond=CARTE, nom="IA")
    texte(d, MARGE + 0.4, iy, gw - 0.8, ih, [
        {"runs": [("L'IA (CLAUDE CODE), POUR DEUX CHOSES", st(ETIQUETTE, espacement=150))], "apres": 6},
        {"runs": [("Générer des scripts.", st(CORPS, taille=15))], "puce": True, "apres": 3},
        {"runs": [("Écrire des guides pas à pas pour monter chaque fonctionnalité dans Unity.",
                   st(CORPS, taille=15))], "puce": True},
    ], ancre="milieu", nom="Texte IA")
    rx = MARGE + gw + 0.5
    rw = L - MARGE - rx
    rh = rw * 9 / 16
    capture(d, "unity", rx, 1.6, rw, rh)
    stats = [("215", "scripts C#"), ("64 000", "lignes de code"), ("23", "shaders faits pour le jeu")]
    sw = rw / 3
    for i, (n, leg) in enumerate(stats):
        texte(d, rx + i * sw, iy + 0.1, sw, ih - 0.1, [
            {"runs": [(n, st(CHIFFRE, taille=34))], "align": "c", "apres": 0},
            {"runs": [(leg, st(LEGENDE, taille=13, couleur=LABEL))], "align": "c"},
        ], nom=f"Chiffre {i + 1}")
    notes(d, "ANTONIN — environ 1 min 15.\n"
             "Le jeu tourne sous Unity 6, avec le pipeline URP, le XR Interaction Toolkit et OpenXR, sur un Meta "
             "Quest 3 relié au PC par Link.\n"
             "Quelques choix nous ont fait gagner du temps. Tous les réglages (arcs, ennemis, vagues, boutique) sont "
             "des ScriptableObjects : on équilibre sans toucher au code. Les flèches, les ennemis et les effets "
             "passent par des pools d'objets, pour garder 72 images par seconde. Les ennemis au sol suivent les voies "
             "grâce au NavMesh. Et un mode clavier-souris permet de tester sans mettre le casque.\n"
             "Côté outils, on a utilisé l'IA, Claude Code, pour deux choses : générer des scripts, et écrire des "
             "guides pas à pas pour monter chaque fonctionnalité dans Unity.\n"
             "Au total : environ 215 scripts, 64 000 lignes de C# et 23 shaders faits pour le jeu.")


def d08_difficultes(deck):
    d = deck.diapo("Ce qui nous a résisté", "Antonin")
    lignes = [
        ("Le mal des transports quand on se déplace",
         "Courir en balançant les bras, rotation par crans, vignette de confort."),
        ("Garder 72 images/s avec des dizaines d'ennemis",
         "Pools d'objets ; ombres et effets réglés au casque."),
        ("13 octobre : le PC tombe en panne, le travail pas poussé est perdu",
         "Projet récupéré sur GitHub, boss final refait. Depuis, on pousse souvent."),
        ("Le casque se figeait en pleine partie",
         "Un bug du pilote NVIDIA : retour à une version plus ancienne."),
        ("Un geste VR qui marche vraiment : le katana",
         "Refait trois fois, testé au casque à chaque version."),
    ]
    mx = 6.62
    texte(d, MARGE + 0.4, 1.5, 5.5, 0.35, [("LE PROBLÈME", st(ETIQUETTE, couleur=DANGER, espacement=200))],
          nom="En-tête problème")
    texte(d, mx + 0.45, 1.5, 5.5, 0.35, [("CE QU'ON A FAIT", st(ETIQUETTE, couleur=VERT, espacement=200))],
          nom="En-tête solution")
    fleche = image_icone("arrowRight", OR)
    y0, rh, gap = 1.95, 0.82, 0.13
    for i, (pb, sol) in enumerate(lignes):
        y = y0 + i * (rh + gap)
        cadre(d, MARGE, y, L - 2 * MARGE, rh, "bouton", fond=CARTE, nom=f"Ligne {i + 1}")
        texte(d, MARGE + 0.4, y, mx - MARGE - 0.75, rh, [(pb, st(CORPS, taille=15, gras=True))], ancre="milieu",
              nom=f"Problème {i + 1}")
        image(d, fleche, mx - 0.05, y + rh / 2 - 0.12, 0.24, 0.24, nom=f"Flèche {i + 1}")
        texte(d, mx + 0.45, y, L - MARGE - mx - 0.8, rh, [(sol, st(CORPS, taille=15))], ancre="milieu",
              nom=f"Solution {i + 1}")
    notes(d, "ANTONIN — environ 1 min 15.\n"
             "Plusieurs choses nous ont résisté. D'abord le confort : se déplacer en VR donne vite mal au cœur. On "
             "court en balançant les bras, avec une rotation par crans et une vignette de confort.\n"
             "Ensuite les performances, avec des dizaines d'ennemis à l'écran : il faut tenir 72 images par seconde.\n"
             "Le 13 octobre, mon PC est tombé en panne : tout ce qui n'était pas poussé sur GitHub était perdu, dont "
             "le boss final, qu'il a fallu refaire. Depuis, on pousse souvent.\n"
             "Le casque se figeait aussi en pleine partie : c'était un bug du pilote NVIDIA, réglé en revenant à une "
             "version plus ancienne.\n"
             "Enfin, un geste en VR doit marcher vraiment dans la main : le katana a été refait trois fois.\n"
             "Dovin va vous parler de l'équipe.")


def d09_repartition(deck):
    d = deck.diapo("Qui a fait quoi", "Dovin")
    membres = [
        ("Antonin", "Antonin Croquette", "Lead developer",
         "Le GDD, le code du jeu, le montage des scènes, l'intégration des branches de l'équipe."),
        ("Jude", "Jude Aybalen", "Scope de départ, en partie · les classes",
         "Le menu des armes, les succès qui débloquent les arcs, les arcs court, puissant et runique."),
        ("Dovin", "Tchoffo Manfouo Dovin Vilson", "Scope de départ, en partie · la grenade · les tests",
         "La grenade à flèches : ses scripts, son prefab, sa ceinture, son affichage sur la montre."),
        ("Merleau", "Merleau Gezo", "Scope de départ, en partie · les tests",
         "Tester le jeu et faire remonter les bugs."),
    ]
    icones = {"Antonin": "gear", "Jude": "trophy", "Dovin": "star", "Merleau": "checkmark"}
    gap = 0.35
    cw = (L - 2 * MARGE - gap) / 2
    ch = (7.0 - 1.55 - 0.2 - gap) / 2
    for i, (o, nom, role, detail) in enumerate(membres):
        x = MARGE + (i % 2) * (cw + gap)
        y = 1.55 + (i // 2) * (ch + gap)
        cadre(d, x, y, cw, ch, nom=f"Membre {i + 1}")
        image(d, image_icone(icones[o], ORATEURS[o]), x + 0.4, y + 0.42, 0.42, 0.42, nom=f"Icône {o}")
        texte(d, x + 1.05, y + 0.32, cw - 1.45, ch - 0.55, [
            {"runs": [(nom, st(CORPS, taille=19, gras=True, couleur=ORATEURS[o]))], "apres": 2},
            {"runs": [(role.upper(), st(ETIQUETTE, taille=11, espacement=120))], "apres": 8},
            {"runs": [(detail, st(CORPS, taille=15))]},
        ], nom=f"Texte {o}")
    notes(d, "DOVIN — environ 45 s.\n"
             "Voici qui a fait quoi. Antonin est le lead developer : le GDD, le code du jeu, le montage des scènes, "
             "et l'intégration de nos branches.\n"
             "Jude, Merleau et moi avons participé au scope de départ. Ensuite, Jude a fait les classes d'arcs et "
             "les succès. Moi, la grenade à flèches : ses scripts, son prefab, sa ceinture et son affichage sur la "
             "montre ; et j'ai testé le jeu. Merleau s'est occupé des tests.")


def d10_collaboration(deck):
    d = deck.diapo("Travailler à quatre", "Dovin")
    # Le graphe des branches, à l'échelle des dates (2 au 23 octobre) : main au milieu, les branches courtes au-dessus,
    # le fork de la grenade en dessous.
    gx0, gx1, my = MARGE + 0.5, L - MARGE - 0.5, 2.6
    haut, bas = 1.98, 3.22

    def x_de(jour):
        return gx0 + (jour - 2) / 21 * (gx1 - gx0)

    def trait(x1, y1, x2, y2, couleur, epaisseur=3.0, nom=None):
        c = d.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = rgb(couleur)
        c.line.width = Pt(epaisseur)
        if nom:
            c.name = nom
        return c

    def point(x, y, couleur, r=0.09, nom=None):
        s = d.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - r), Inches(y - r), Inches(2 * r), Inches(2 * r))
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(couleur)
        s.line.color.rgb = rgb(FOND)
        s.line.width = Pt(1.5)
        if nom:
            s.name = nom
        return s

    def branche(debut, fin, by, couleur, nom):
        """Une branche qui part de main et y revient ; les branches d'un jour gardent un palier lisible."""
        xa, xb = x_de(debut), x_de(fin)
        biais = 0.28
        xa = min(xa, xb - 2 * biais - 0.35)
        trait(xa, my, xa + biais, by, couleur, nom=f"{nom} départ")
        trait(xa + biais, by, xb - biais, by, couleur, nom=f"{nom} branche")
        trait(xb - biais, by, xb, my, couleur, nom=f"{nom} fusion")
        return (xa + xb) / 2

    centre_td = branche(11, 12, haut, ORATEURS["Antonin"], "Tower defense")
    centre_classes = branche(18, 19, haut, ORATEURS["Jude"], "Classes")
    branche(14, 22, bas, ORATEURS["Dovin"], "Grenade")
    trait(gx0, my, gx1, my, OR, 3.5, nom="main")
    for j in (18, 21):  # la branche de la grenade reprend main
        trait(x_de(j), my, x_de(j) + 0.14, bas, ORATEURS["Dovin"], 1.25, nom=f"Mise à jour grenade {j}")
    for j, c in ((2, OR), (12, ORATEURS["Antonin"]), (13, DANGER), (16.5, ORATEURS["Merleau"]),
                 (19, ORATEURS["Jude"]), (22, ORATEURS["Dovin"]), (23, OR)):
        point(x_de(j), my, c, nom=f"Repère {j}")

    def etiquette(t, x, y, w, couleur, align, style=None):
        texte(d, x, y, w, 0.3, [{"runs": [(t, style or st(ETIQUETTE, taille=11, couleur=couleur, espacement=100))],
                                 "align": align}], nom=f"Étiquette {t}")

    # Au-dessus de main : la panne (vers la droite) et la démo (vers la gauche), entre les branches.
    etiquette("13 OCT · PANNE DU PC", x_de(13) - 0.08, my - 0.4, 2.3, DANGER, "g")
    etiquette("23 OCT · DÉMO", x_de(23) - 1.9, my - 0.4, 2.0, OR, "d")
    # En dessous : le début et la Nuit des assos, entre les diagonales du fork.
    etiquette("2 OCT · DÉBUT", x_de(2) - 0.08, my + 0.16, 2.3, OR, "g")
    etiquette("NUIT DES ASSOS", x_de(16.5) - 0.8, my + 0.16, 1.6, ORATEURS["Merleau"], "c")
    etiquette("tower defense · pull request", centre_td - 1.4, haut - 0.38, 2.8, None, "c",
              st(LEGENDE, couleur=ORATEURS["Antonin"]))
    etiquette("classes-arc · Jude · 19 oct", centre_classes - 1.4, haut - 0.38, 2.8, None, "c",
              st(LEGENDE, couleur=ORATEURS["Jude"]))
    etiquette("fork de Dovin · grenade-fleches · fusionné le 22 oct", x_de(18) - 2.3, bas + 0.1, 4.6, None, "c",
              st(LEGENDE, couleur=ORATEURS["Dovin"]))
    # En bas : s'organiser, tester, et la capture GitHub.
    by = 4.25
    col = [
        ("S'ORGANISER", [
            "Un dépôt GitHub ; main reste toujours jouable.",
            "Une branche par fonctionnalité, fusionnée par pull request.",
            "Le GDD, mis à jour et daté à chaque décision.",
        ]),
        ("TESTER", [
            "Au casque, après chaque ajout.",
            "À la Nuit des assos, avec de vrais joueurs : vagues 3 à 5 plus dures, Facile plus doux, dragon "
            "renforcé.",
        ]),
    ]
    cw = 3.95
    for i, (titre, items) in enumerate(col):
        x = MARGE + i * (cw + 0.35)
        paras = [{"runs": [(titre, st(ETIQUETTE, espacement=200))], "apres": 6}]
        for it in items:
            paras.append({"runs": [(it, st(CORPS, taille=15))], "puce": True, "apres": 5})
        texte(d, x, by, cw, 2.6, paras, nom=f"Colonne {titre}")
    gx = MARGE + 2 * (cw + 0.35) + 0.1
    gw = L - MARGE - gx
    capture(d, "github", gx, by + 0.05, gw, gw * 9 / 16)
    notes(d, "DOVIN — environ 1 min.\n"
             "Pour travailler ensemble, on a un dépôt GitHub. La branche main reste toujours jouable, et chaque "
             "fonctionnalité a sa branche : le tower defense est passé par une pull request, les classes de Jude "
             "par la branche classes-arc, et ma grenade par un fork, fusionné le 22 octobre.\n"
             "Le GDD sert de référence commune : on le met à jour et on date chaque décision.\n"
             "Pour les tests : au casque après chaque ajout, et à la Nuit des assos, avec de vrais joueurs. Leurs "
             "retours ont changé l'équilibrage : les vagues 3 à 5 étaient trop faciles, le mode Facile pas assez "
             "doux, et le dragon trop faible.")


def d11_bilan(deck):
    d = deck.diapo("Bilan", "Dovin")
    chiffres = [("6", "vagues, 5 boss et MK"), ("8", "armes à jouer"), ("23", "améliorations"),
                ("2", "secrets sur la carte")]
    lw = 4.9
    cadre(d, MARGE, 1.55, lw, 5.15, nom="Chiffres")
    texte(d, MARGE + 0.4, 1.95, lw - 0.8, 0.75, [
        {"runs": [("TOUTES LES CONSIGNES DU COURS", st(ETIQUETTE, espacement=150))], "align": "c", "apres": 2},
        {"runs": [("sont dans le jeu, et en plus :", st(CORPS, taille=15, couleur=LABEL))], "align": "c"},
    ], nom="Consignes")
    for i, (n, leg) in enumerate(chiffres):
        x = MARGE + 0.3 + (i % 2) * (lw - 0.6) / 2
        y = 2.95 + (i // 2) * 1.75
        texte(d, x, y, (lw - 0.6) / 2, 1.5, [
            {"runs": [(n, st(CHIFFRE, taille=46))], "align": "c", "apres": 0},
            {"runs": [(leg, st(CORPS, taille=15, couleur=LABEL))], "align": "c"},
        ], ancre="milieu", nom=f"Chiffre {i + 1}")
    rx = MARGE + lw + 0.5
    rw = L - MARGE - rx
    blocs = [
        ("CE QU'ON RETIENT", [
            "En VR, rien n'est fini tant qu'on ne l'a pas essayé au casque.",
            "Des joueurs extérieurs voient ce qu'on ne voit plus.",
            "Pousser son travail souvent : on l'a appris le 13 octobre.",
        ]),
        ("ET APRÈS", [
            "Une version autonome sur le Quest, sans PC.",
            "Une nouvelle arène pour MK, et une histoire cachée, d'easter egg en easter egg.",
            "De vraies musiques et de vrais sons.",
        ]),
    ]
    paras = []
    for i, (titre, items) in enumerate(blocs):
        paras.append({"runs": [(titre, st(ETIQUETTE, espacement=200))], "apres": 6, "avant": 0 if i == 0 else 22})
        for it in items:
            paras.append({"runs": [(it, st(CORPS, taille=17))], "puce": True, "apres": 6})
    texte(d, rx, 1.6, rw, 5.0, paras, ancre="milieu", nom="Retenir et après")
    notes(d, "DOVIN — environ 45 s.\n"
             "En bilan : toutes les consignes du cours sont dans le jeu, avec six vagues, cinq boss et un dragon, "
             "huit armes et deux secrets.\n"
             "On retient trois choses : en VR, rien n'est fini tant qu'on ne l'a pas essayé au casque ; des joueurs "
             "extérieurs voient ce qu'on ne voit plus ; et il faut pousser son travail souvent.\n"
             "Pour la suite : une version autonome sur le Quest, sans PC ; une nouvelle arène pour MK et une histoire "
             "cachée, révélée d'easter egg en easter egg ; et de vraies musiques.\n"
             "Place à la démo.")


def d12_demo(deck):
    d = deck.diapo(fond=True)
    plein = capture_plein_ecran(d, "mk", voile_gauche=0.85, voile_bas=0.3)
    texte(d, 0.85, 2.15, 8, 0.4, [("PLACE AU JEU", st(ETIQUETTE, taille=14, espacement=300))], nom="Accroche")
    texte(d, 0.8, 2.5, 9, 1.3, [("La démo", st(TITRE, taille=66))], ancre="milieu", nom="Titre démo")
    texte(d, 0.85, 3.95, 7.6, 1.6, [
        {"runs": [(liste("Un tir parfait", "une vague et son boss", "construire au marteau", "MK, le dragon"),
                   st(CORPS, taille=20))], "apres": 14},
        {"runs": [("Puis vos questions.", st(CORPS, taille=18, couleur=LABEL))]},
    ], nom="Programme démo")
    if not plein:
        capture(d, "mk", 8.75, 1.9, 3.98, 3.98 * 9 / 16)
    notes(d, "DÉMO — 5 min. Celui qui joue commente lentement ; un autre tape les raccourcis au clavier "
             "(liste dans Docs/Demo_Vendredi.md, partie 4).\n"
             "Ordre : un tir parfait aux cibles d'entraînement ; le gong, une vague (récolter, construire une "
             "baliste au marteau) ; le boss et son point faible ; puis MK (F, puis N pour avancer).")


def generer(chemin):
    deck = Deck()
    for f in (d01_titre, d02_programme, d03_concept, d04_partie, d05_consignes, d06_ajouts, d07_technique,
              d08_difficultes, d09_repartition, d10_collaboration, d11_bilan, d12_demo):
        f(deck)
    assert deck.numero == deck.total, deck.numero
    deck.prs.core_properties.title = "THUNK VR"
    deck.prs.core_properties.author = "Antonin Croquette, Jude Aybalen, Tchoffo Manfouo Dovin Vilson, Merleau Gezo"
    deck.prs.save(chemin)


# ---------------------------------------------------------------------------------------------------------------
# PowerPoint : polices intégrées, PDF, aperçus.

POWERSHELL = r"""
$ErrorActionPreference = 'Stop'
$app = New-Object -ComObject PowerPoint.Application
$pres = $app.Presentations.Open('{source}', -1, 0, 0)
try {{
  $pres.SaveAs('{pptx}', 24, -1)
  $pres.SaveAs('{pdf}', 32)
  if ('{apercus}') {{
    foreach ($s in $pres.Slides) {{ $s.Export(('{apercus}\diapo-{{0:D2}}.png' -f $s.SlideIndex), 'PNG', 1600, 900) }}
  }}
}} finally {{
  $pres.Close()
  if ($app.Presentations.Count -eq 0) {{ $app.Quit() }}
}}
"""


def finir_avec_powerpoint(brut, apercus):
    script = POWERSHELL.format(source=brut, pptx=SORTIE, pdf=SORTIE.with_suffix(".pdf"), apercus=apercus or "")
    r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", script],
                       capture_output=True, text=True)
    if r.returncode != 0:
        sys.exit("PowerPoint n'a pas pu finir la présentation :\n" + r.stderr + r.stdout)


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--apercus", help="dossier où écrire une image PNG par diapo")
    args = parser.parse_args()
    brut = TEMP / "brut.pptx"
    generer(brut)
    apercus = None
    if args.apercus:
        apercus = Path(args.apercus).resolve()
        apercus.mkdir(parents=True, exist_ok=True)
    finir_avec_powerpoint(brut, apercus)
    manquantes = [n for n in CAPTURES_ATTENDUES if not chercher_capture(n)]
    print(f"Présentation : {SORTIE}")
    print(f"PDF de secours : {SORTIE.with_suffix('.pdf')}")
    if manquantes:
        print("Captures manquantes (dans Presentation/captures/) : " + ", ".join(manquantes))


if __name__ == "__main__":
    main()
